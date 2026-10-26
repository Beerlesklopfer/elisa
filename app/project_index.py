"""
Indexierung der Projektdokumente (RAG, projekt-skopiert) über Postgres+pgvector.

Jedes Dokument eines Projekts wird eingebettet und dem Projekt zugeordnet
(`EmbeddingChunk.collection = "proj-<key>"`, plus `project`-FK). Der Chat zieht
sie beratend als Kontext heran. Kein Fine-Tuning; die DB bleibt Quelle der
Wahrheit (CLAUDE.md).
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from django.conf import settings

from app import rag_service

CHUNK_SIZE = 800
CHUNK_OVERLAP = 150


def collection_name(project) -> str:
    key = re.sub(r"[^a-zA-Z0-9_-]", "-", project.key or str(project.pk))
    return f"proj-{key}"[:64]


def _chunks(text: str):
    out, start = [], 0
    while start < len(text):
        out.append(text[start:start + CHUNK_SIZE])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return out


def extract_text(path: Path):
    """Liefert [(seite, text)] für PDF/DOCX/Text. Tolerant bei Fehlern."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            return [(i + 1, (p.extract_text() or ""))
                    for i, p in enumerate(PdfReader(str(path)).pages)]
        if suffix in (".docx", ".odt", ".doc"):
            out = subprocess.run(
                [settings.PANDOC_BIN, str(path), "-t", "plain"],
                capture_output=True, text=True, check=True).stdout
            return [(1, out)]
        if suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            pages = []
            for i, ws in enumerate(wb.worksheets):
                rows = ["\t".join("" if c is None else str(c) for c in row)
                        for row in ws.iter_rows(values_only=True)]
                pages.append((i + 1, f"[{ws.title}]\n" + "\n".join(rows)))
            return pages or [(1, "")]
        if suffix == ".pptx":
            from pptx import Presentation
            return [(i + 1, "\n".join(sh.text for sh in slide.shapes
                                      if getattr(sh, "text", "")))
                    for i, slide in enumerate(Presentation(str(path)).slides)]
        return [(1, path.read_text(errors="ignore"))]
    except Exception:
        return []


def index_document(doc) -> int:
    """Indexiert ein ProjectDocument in die Projekt-Collection. Liefert #Chunks."""
    from app.models import EmbeddingChunk
    path = Path(doc.file.path)
    if not path.is_file():
        return 0
    coll = collection_name(doc.project)
    name = doc.original_name or path.name   # Originalname als RAG-Quelle (Pfad ist UUID)
    # SHA-256 der Datei (Inhalts-Hash) am Dokument festhalten.
    file_hash = rag_service.file_sha256(path)
    if doc.content_hash != file_hash:
        doc.content_hash = file_hash
        doc.save(update_fields=["content_hash"])
    # Alte Chunks dieses Dokuments ersetzen (Upsert-Semantik).
    EmbeddingChunk.objects.filter(collection=coll, source=name).delete()
    rows, n = [], 0
    for page_no, text in extract_text(path):
        text = (text or "").strip()
        if not text:
            continue
        chunks = _chunks(text)
        if not chunks:
            continue
        vecs = rag_service.embed_many(chunks)
        for j, (chunk, vec) in enumerate(zip(chunks, vecs)):
            rows.append(EmbeddingChunk(
                collection=coll, project=doc.project, source=name,
                kind=doc.kind, page=page_no, text=chunk, content_hash=file_hash,
                uid=rag_service.make_uid(coll, doc.pk, name, page_no, j),
                embedding=vec))
        n += len(chunks)
    if rows:
        EmbeddingChunk.objects.bulk_create(rows, batch_size=200)
    # EVENT (Upload): Esra soll die neuen Chunks verschlagworten.
    if n:
        from app.models import enqueue_esra, EsraTask
        enqueue_esra(doc.project, EsraTask.Kind.TAG_DOCUMENT, document=doc)
    return n


def remove_document(doc) -> None:
    """Entfernt die Chunks eines Dokuments aus der Projekt-Collection."""
    from app.models import EmbeddingChunk
    try:
        EmbeddingChunk.objects.filter(
            collection=collection_name(doc.project),
            source=doc.original_name or Path(doc.file.name).name).delete()
    except Exception:
        pass


def retrieve_project_context(project, question: str, n: int = 4):
    """RAG-Treffer aus der Projekt-Collection (oder [] wenn leer/Fehler)."""
    try:
        return rag_service._search(collection_name(project), question, n)
    except Exception:
        return []
